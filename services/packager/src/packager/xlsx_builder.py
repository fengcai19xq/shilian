"""核对表：清单项 / 对应文件 / 页码 / 状态 / 责任部门 / 备注。

缺失项整行标红并注明责任部门；待确认项标黄。页码为合订 PDF 中的物理页码区间。
文档属性里的创建/修改时间属于允许的重放差异。
"""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .layout import Layout
from .models import EntryStatus, Manifest

HEADERS = ["编号", "清单项", "对应文件", "页码", "状态", "责任部门", "备注"]
STATUS_LABEL = {
    EntryStatus.MATCHED: "已匹配",
    EntryStatus.PENDING: "待确认",
    EntryStatus.MISSING: "缺失",
}
MISSING_FILL = PatternFill("solid", start_color="FFC7CE", end_color="FFC7CE")
MISSING_FONT = Font(color="9C0006", bold=True)
PENDING_FILL = PatternFill("solid", start_color="FFEB9C", end_color="FFEB9C")
UNASSIGNED_DEPT = "待指派"


def build_checklist_xlsx(
    manifest: Manifest, layout: Layout, members: dict[str, list[str]]
) -> bytes:
    wb = Workbook()
    wb.properties.creator = "shilian-packager"
    ws = wb.active
    ws.title = "核对表"
    ws.append([manifest.title])
    ws.append([f"资料包编号：{manifest.package_id}"])
    ws.append(HEADERS)
    header_row = ws.max_row
    for cell in ws[header_row]:
        cell.font = Font(bold=True)

    for entry in manifest.entries:
        span = layout.span_of(entry.no)
        files = "\n".join(members.get(entry.no, [])) or "—"
        dept = (entry.owner_dept or UNASSIGNED_DEPT) if entry.is_missing else ""
        ws.append(
            [
                entry.no,
                entry.std_name,
                files,
                span.label if span else "—",
                STATUS_LABEL[entry.status],
                dept,
                entry.note,
            ]
        )
        row = ws[ws.max_row]
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if entry.is_missing:
                cell.fill = MISSING_FILL
                cell.font = MISSING_FONT
            elif entry.status is EntryStatus.PENDING:
                cell.fill = PENDING_FILL

    for col, width in zip("ABCDEFG", (8, 28, 48, 10, 10, 14, 36)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
