"""执行器端到端测试：合订、水印、目录页码、缺失项、重放一致性、异常分支。"""

from __future__ import annotations

import io
import re
import zipfile

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from packager import (
    LocalStorage,
    OutputKind,
    PackagingError,
    Storage,
    StorageError,
    build_artifacts,
    execute,
)
from packager.layout import TOC_ROWS_PER_PAGE

WATERMARK = "仅供示例银行授信审查使用 · 2026-03-12"
# 封面 1 + 目录 1 + 营业执照 3 + 审计报告 5+2 + 银行流水 2
EXPECTED_PAGES = 1 + 1 + 3 + 7 + 2


def _pdf_texts(data: bytes) -> list[str]:
    return [p.extract_text() for p in PdfReader(io.BytesIO(data)).pages]


def _rows(data: bytes) -> list[tuple]:
    ws = load_workbook(io.BytesIO(data)).active
    return [tuple(c.value for c in row) for row in ws.iter_rows()]


@pytest.fixture
def result(manifest, storage):
    return execute(manifest, storage)


def _read_artifact(result, kind: OutputKind) -> bytes:
    art = result.artifact(kind)
    assert art is not None
    with open(art.uri, "rb") as fh:
        return fh.read()


# ---------- 执行结果 ----------


def test_execute_writes_all_outputs(result, storage):
    assert isinstance(storage, Storage)
    assert [a.kind for a in result.artifacts] == list(OutputKind)
    assert result.requires_manual_review is True
    assert result.missing_entries == ["6.3", "8"]
    for art in result.artifacts:
        assert art.uri.startswith(str(storage.out_dir))
        assert art.size > 0
        assert storage.signed_url(art.uri).startswith("file://")


def test_outputs_subset_and_dedup(manifest, storage):
    manifest["outputs"] = ["checklist_xlsx", "checklist_xlsx"]
    res = execute(manifest, storage)
    assert [a.kind for a in res.artifacts] == [OutputKind.CHECKLIST_XLSX]


# ---------- 合订 PDF ----------


def test_merged_pdf_order_and_page_count(result):
    texts = _pdf_texts(_read_artifact(result, OutputKind.MERGED_PDF))
    assert len(texts) == EXPECTED_PAGES
    assert "示例银行授信资料包" in texts[0]
    assert "人工终审" in texts[0]
    body = [re.search(r"([A-Z0-9]+ page \d+)", t).group(1) for t in texts[2:]]
    assert body == [
        "LICENSE page 1",
        "LICENSE page 2",
        "LICENSE page 3",
        # order=1 的 audit.pdf 排在 order=2 的 audit_2.pdf 前面
        *[f"AUDIT page {i}" for i in range(1, 6)],
        "AUDIT2 page 1",
        "AUDIT2 page 2",
        # pages=[2,3] 只取第 2、3 页
        "BANK page 2",
        "BANK page 3",
    ]


def test_watermark_and_footer_on_every_page(result):
    texts = _pdf_texts(_read_artifact(result, OutputKind.MERGED_PDF))
    for idx, text in enumerate(texts, start=1):
        assert WATERMARK in text, f"第 {idx} 页缺水印"
        assert f"- {idx} / {len(texts)} -" in text


def test_watermark_falls_back_to_title(manifest, storage):
    manifest["watermark"] = ""
    pdf = build_artifacts(manifest, storage.read)[OutputKind.MERGED_PDF]
    for text in _pdf_texts(pdf):
        assert text.count(manifest["title"]) >= 3


def test_toc_page_numbers_point_to_entry_first_page(result):
    texts = _pdf_texts(_read_artifact(result, OutputKind.MERGED_PDF))
    toc = texts[1]
    assert "目" in toc
    expected = {"营业执照": ("LICENSE page 1", 3), "最近三年审计报告": ("AUDIT page 1", 6)}
    expected["银行流水"] = ("BANK page 2", 13)
    for name, (first_line, page_no) in expected.items():
        assert re.search(rf"{name}\s*\n{page_no}\n", toc), f"目录缺 {name} → {page_no}"
        assert first_line in texts[page_no - 1]
    assert "公司章程" not in toc
    assert "纳税证明" not in toc


def test_toc_spans_multiple_pages(storage):
    count = TOC_ROWS_PER_PAGE + 2
    m = {
        "package_id": "pkg_many",
        "title": "多条目资料包",
        "watermark": "WM",
        "outputs": ["merged_pdf"],
        "entries": [
            {"no": str(i), "std_name": f"条目{i}", "files": [{"uri": "docs/license.pdf"}]}
            for i in range(1, count + 1)
        ],
    }
    texts = _pdf_texts(build_artifacts(m, storage.read)[OutputKind.MERGED_PDF])
    assert len(texts) == 1 + 2 + 3 * count
    assert "目录（续）" in texts[2]
    # 第一个条目从第 4 页开始（封面 1 + 目录 2），最后一个条目在续页上
    assert re.search(r"条目1\s*\n4\n", texts[1])
    last_start = 4 + 3 * (count - 1)
    assert re.search(rf"条目{count}\s*\n{last_start}\n", texts[2])
    assert "LICENSE page 1" in texts[last_start - 1]


# ---------- zip ----------


def test_zip_layout_and_names(result):
    with zipfile.ZipFile(io.BytesIO(_read_artifact(result, OutputKind.ZIP))) as zf:
        names = zf.namelist()
        assert names == [
            "01_营业执照/01_营业执照.pdf",
            "03.1_最近三年审计报告/03.1_最近三年审计报告_2021-2023_1.pdf",
            "03.1_最近三年审计报告/03.1_最近三年审计报告_2021-2023_2.pdf",
            "07_银行流水/07_银行流水.pdf",
        ]
        bank = PdfReader(io.BytesIO(zf.read("07_银行流水/07_银行流水.pdf")))
        assert [p.extract_text().strip() for p in bank.pages] == ["BANK page 2", "BANK page 3"]
        assert all(i.date_time == (1980, 1, 1, 0, 0, 0) for i in zf.infolist())


# ---------- 核对表 ----------


def test_checklist_rows_and_missing_highlight(result):
    data = _read_artifact(result, OutputKind.CHECKLIST_XLSX)
    ws = load_workbook(io.BytesIO(data)).active
    rows = _rows(data)
    assert rows[2] == ("编号", "清单项", "对应文件", "页码", "状态", "责任部门", "备注")
    by_no = {r[0]: r for r in rows[3:]}
    assert list(by_no) == ["1", "3.1", "6.3", "7", "8"]

    assert by_no["1"][3] == "3-5"
    assert by_no["3.1"][3] == "6-12"
    assert by_no["7"][3:5] == ("13-14", "待确认")
    assert by_no["3.1"][2].count("\n") == 1

    missing = by_no["6.3"]
    assert missing[2:7] == ("—", "—", "缺失", "法务部", "库内仅 2022 版，需最新备案版本")
    assert by_no["8"][5] == "待指派"

    for row in ws.iter_rows(min_row=4):
        no = row[0].value
        red = row[0].fill.start_color.rgb.endswith("FFC7CE")
        assert red == (no in {"6.3", "8"}), f"{no} 标红状态不对"
        if red:
            assert all(c.font.color.rgb.endswith("9C0006") for c in row)


def test_missing_entries_never_read(manifest, storage):
    """missing 条目即使带了 files 也不读取、不入 zip/pdf。"""
    read_log: list[str] = []

    def spy(uri: str) -> bytes:
        read_log.append(uri)
        return storage.read(uri)

    build_artifacts(manifest, spy)
    assert "oss://docs/should_not_be_read.pdf" not in read_log
    # 同一 uri 只读一次
    assert len(read_log) == len(set(read_log)) == 4


# ---------- 重放一致性 ----------


def test_replay_is_identical(manifest, storage, tmp_path):
    first = build_artifacts(manifest, storage.read)
    other = LocalStorage(storage.root, out_dir=tmp_path / "out2")
    second = build_artifacts(manifest, other.read)

    assert first[OutputKind.ZIP] == second[OutputKind.ZIP]
    assert first[OutputKind.MERGED_PDF] == second[OutputKind.MERGED_PDF]
    # xlsx 容器内含保存时间，按内容比较
    assert _rows(first[OutputKind.CHECKLIST_XLSX]) == _rows(second[OutputKind.CHECKLIST_XLSX])


def test_execute_replay_overwrites_same_uri(manifest, storage):
    a = execute(manifest, storage)
    b = execute(manifest, storage)
    assert [x.uri for x in a.artifacts] == [x.uri for x in b.artifacts]


# ---------- 异常分支 ----------


def test_page_out_of_range(manifest, storage):
    manifest["entries"][0]["files"][0]["pages"] = [4]
    with pytest.raises(PackagingError, match="页码越界"):
        build_artifacts(manifest, storage.read)


def test_non_pdf_source(manifest, storage):
    manifest["entries"][0]["files"][0]["uri"] = "docs/note.txt"
    with pytest.raises(PackagingError, match="不是可解析的 PDF"):
        build_artifacts(manifest, storage.read)


def test_missing_source_file(manifest, storage):
    manifest["entries"][0]["files"][0]["uri"] = "docs/nope.pdf"
    with pytest.raises(StorageError):
        execute(manifest, storage)


def test_invalid_manifest():
    with pytest.raises(ValueError):
        build_artifacts({"package_id": "x", "title": "t", "outputs": ["docx"]}, lambda u: b"")
    with pytest.raises(ValueError):
        build_artifacts(
            {
                "package_id": "x",
                "title": "t",
                "entries": [{"no": "1", "std_name": "a", "files": [{"uri": "a", "pages": [0]}]}],
            },
            lambda u: b"",
        )
