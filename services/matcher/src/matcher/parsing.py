"""清单解析：Excel（openpyxl）或纯文本 -> 结构化 ChecklistItem。"""

from __future__ import annotations

import io
import re
import uuid
from pathlib import Path

import yaml
from openpyxl import load_workbook

from .config import DATA_DIR, ParserConfig
from .extract import extract_constraints
from .models import Checklist, ChecklistItem
from .typedict import TypeDictionary, _normalize, default_dictionary

_TEXT_LINE = re.compile(r"^\s*(?P<no>[0-9]+(?:\.[0-9]+)*)[、.．:：)\s]+\s*(?P<rest>.+?)\s*$")


def load_column_aliases(path: Path | str | None = None) -> dict[str, list[str]]:
    path = Path(path) if path else DATA_DIR / "columns.yaml"
    data = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    return data.get("columns") or {}


class ChecklistParser:
    def __init__(
        self,
        config: ParserConfig | None = None,
        dictionary: TypeDictionary | None = None,
    ):
        self.config = config or ParserConfig()
        self.dictionary = dictionary or default_dictionary()
        aliases = self.config.column_aliases or load_column_aliases()
        # 归一后的列头 -> 字段名
        self._header_index = {
            _normalize(alias): field for field, alias_list in aliases.items() for alias in alias_list
        }

    # ---------- 公共入口 ----------

    def parse(self, content: bytes | str, filename: str = "", title: str = "") -> Checklist:
        """按文件名后缀分派到 Excel / 文本解析。"""
        if filename.lower().endswith((".xlsx", ".xlsm")):
            assert isinstance(content, bytes), "Excel 清单需要传字节流"
            return self.parse_excel(content, title=title or Path(filename).stem)
        text = content.decode("utf-8") if isinstance(content, bytes) else content
        return self.parse_text(text, title=title or (Path(filename).stem if filename else "清单"))

    def parse_excel(self, content: bytes, title: str = "清单") -> Checklist:
        wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
        ws = wb.active
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()

        header_row, mapping = self._locate_header(rows)
        checklist = self._new_checklist(title)
        if mapping is None:
            return checklist

        auto_no = 0
        for raw_row in rows[header_row + 1 :]:
            name = self._cell(raw_row, mapping.get("name"))
            requirement = self._cell(raw_row, mapping.get("requirement"))
            no = self._cell(raw_row, mapping.get("no"))
            if not name and not requirement:
                continue
            if not name:
                continue
            if not no:
                auto_no += 1
                no = str(auto_no)
            raw_text = name if not requirement else f"{name}（{requirement}）"
            checklist.items.append(self._build_item(checklist.checklist_id, no, raw_text, name))
        return checklist

    def parse_text(self, text: str, title: str = "清单") -> Checklist:
        checklist = self._new_checklist(title)
        auto_no = 0
        for line in text.splitlines():
            line = line.strip()
            if not line:
                continue
            m = _TEXT_LINE.match(line)
            if m:
                no, raw_text = m.group("no"), m.group("rest")
            else:
                auto_no += 1
                no, raw_text = str(auto_no), line
            checklist.items.append(self._build_item(checklist.checklist_id, no, raw_text, raw_text))
        return checklist

    # ---------- 内部 ----------

    def _new_checklist(self, title: str) -> Checklist:
        return Checklist(checklist_id=f"cl_{uuid.uuid4().hex[:8]}", title=title or "清单")

    def _build_item(self, checklist_id: str, no: str, raw_text: str, name_text: str) -> ChecklistItem:
        hit = self.dictionary.resolve(name_text)
        return ChecklistItem(
            item_id=f"it_{uuid.uuid4().hex[:8]}",
            checklist_id=checklist_id,
            no=str(no).strip(),
            raw_text=raw_text,
            std_type=hit.std_type,
            type_status=hit.status,
            std_name=hit.std_name,
            constraints=extract_constraints(raw_text, self.config.reference_year),
        )

    def _locate_header(self, rows: list[list]) -> tuple[int, dict[str, int] | None]:
        """在前若干行里找列头行，返回行号与「字段 -> 列下标」。"""
        for idx, row in enumerate(rows[: self.config.max_header_scan_rows]):
            mapping: dict[str, int] = {}
            for col, value in enumerate(row):
                field = self._header_index.get(_normalize(str(value or "")))
                if field and field not in mapping:
                    mapping[field] = col
            if "name" in mapping:
                return idx, mapping
        return 0, None

    @staticmethod
    def _cell(row: list, col: int | None) -> str:
        if col is None or col >= len(row):
            return ""
        value = row[col]
        return "" if value is None else str(value).strip()
